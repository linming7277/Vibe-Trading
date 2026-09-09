# -*- coding: utf-8 -*-
"""Export current L3 leaders using the same columns as 当前三级行业龙头池_2026-08-21.xlsx."""
from __future__ import annotations

import json
import sqlite3
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

RESEARCH_DB = Path.home() / ".vibe-trading" / "research.db"
TDX_DB = Path.home() / ".vibe-trading" / "tdx_data.db"
OUT_DIR = Path(r"d:/AI/hzstock/tmp")
OUT_DIR.mkdir(parents=True, exist_ok=True)

HEADER = [
    "序号",
    "股票代码",
    "股票名称",
    "一级行业",
    "二级行业",
    "细分赛道（三级/末级）",
    "赛道内龙头排名",
    "综合评分",
    "评分说明（构成项）",
    "覆盖率",
    "状态",
    "公司简介（本地通达信资料）",
]


def first_text(*values: object) -> str:
    for value in values:
        text = str(value or "").strip()
        if text and text not in {"--", "0", "0.0"}:
            return text
    return ""


def compact_text(value: object, max_length: int = 160) -> str:
    cleaned = " ".join(str(value or "").split()).strip()
    if len(cleaned) > max_length:
        return cleaned[: max_length - 1] + "…"
    return cleaned


def company_introduction(stock_code: str, fundamentals: dict, details: dict, profiles: dict) -> str:
    fundamental = fundamentals.get(stock_code) or {}
    detail = details.get(stock_code) or {}
    profile = profiles.get(stock_code) or {}
    base = fundamental.get("base_raw") or {}
    extended = fundamental.get("extended_raw") or {}
    detail_extended = detail.get("extended") or {}
    description = first_text(
        profile.get("company_description"),
        detail_extended.get("CompanyDescription"),
        detail_extended.get("CompanyIntroduction"),
        detail_extended.get("CompanyIntro"),
        detail_extended.get("GSJJ"),
        extended.get("CompanyDescription"),
        extended.get("CompanyIntroduction"),
        extended.get("CompanyIntro"),
        extended.get("GSJJ"),
        base.get("CompanyDescription"),
        base.get("CompanyIntroduction"),
        base.get("CompanyIntro"),
        base.get("GSJJ"),
    )
    main_business = first_text(
        profile.get("main_business"),
        fundamental.get("main_business"),
        detail_extended.get("MainBusiness"),
        detail_extended.get("ZYYW"),
        extended.get("MainBusiness"),
        extended.get("ZYYW"),
        base.get("MainBusiness"),
        base.get("ZYYW"),
    )
    products = first_text(
        profile.get("main_products"),
        detail_extended.get("MainProducts"),
        detail_extended.get("MainProduct"),
        detail_extended.get("Products"),
        detail_extended.get("ZYCP"),
        extended.get("MainProducts"),
        extended.get("MainProduct"),
        extended.get("Products"),
        extended.get("ZYCP"),
        base.get("MainProducts"),
        base.get("MainProduct"),
        base.get("Products"),
        base.get("ZYCP"),
    )
    business_scope = first_text(
        profile.get("business_scope"),
        detail_extended.get("BusinessScope"),
        detail_extended.get("BusinessRange"),
        detail_extended.get("OperationScope"),
        extended.get("BusinessScope"),
        extended.get("BusinessRange"),
        extended.get("OperationScope"),
        base.get("BusinessScope"),
        base.get("BusinessRange"),
        base.get("OperationScope"),
    )
    if description:
        return compact_text(description)
    if main_business:
        return compact_text(f"主营业务：{main_business}")
    if products:
        return compact_text(f"主要产品：{products}")
    if business_scope:
        return compact_text(f"经营范围：{business_scope}")
    return "本地暂无可靠公司简介"


def score_explanation(row: dict) -> str:
    try:
        score = json.loads(row.get("component_scores_json") or "{}")
    except json.JSONDecodeError:
        score = {}

    def fmt(key: str) -> str:
        value = score.get(key)
        return "—" if value is None else f"{float(value):.1f}"

    total = row.get("leader_score")
    total_text = "—" if total is None else f"{float(total):.2f}"
    return "；".join([
        f"质量分 {total_text}/100",
        f"规模地位 {fmt('industry_position')}",
        f"盈利质量 {fmt('profitability')}",
        f"增长稳定 {fmt('growth_stability')}",
        f"现金流 {fmt('cash_flow')}",
        f"估值 {fmt('valuation')}",
        f"治理风险 {fmt('governance_risk')}",
    ])


def load_payload_map(conn: sqlite3.Connection, dataset: str) -> dict[str, dict]:
    out: dict[str, dict] = {}
    for record_key, payload_json in conn.execute(
        "SELECT record_key, payload_json FROM records WHERE dataset=?",
        (dataset,),
    ):
        try:
            out[str(record_key)] = json.loads(payload_json or "{}")
        except json.JSONDecodeError:
            out[str(record_key)] = {}
    return out


research = sqlite3.connect(str(RESEARCH_DB))
research.row_factory = sqlite3.Row

pool_cols = [c[1] for c in research.execute("PRAGMA table_info(l3_leader_pool_runs)")]
pool = dict(research.execute(
    """SELECT * FROM l3_leader_pool_runs
       WHERE status='COMPLETED'
       ORDER BY as_of DESC, completed_at DESC
       LIMIT 1"""
).fetchone())

# Prefer the pool that low-value currently references if fresher membership exists
active_pool = research.execute(
    """SELECT pool_id, MAX(last_seen_at) AS last_seen
       FROM l3_leader_pool_members
       WHERE lifecycle_status IN ('NEW','ACTIVE','REENTERED')
       GROUP BY pool_id
       ORDER BY last_seen DESC
       LIMIT 1"""
).fetchone()
pool_id = active_pool["pool_id"] if active_pool else pool["id"]
pool_meta = dict(research.execute(
    "SELECT * FROM l3_leader_pool_runs WHERE id=?", (pool_id,)
).fetchone() or pool)
as_of = str(pool_meta.get("as_of") or "")

leaders = [dict(r) for r in research.execute(
    """SELECT stock_code, stock_name, level1_name, level2_name, level3_name,
              leader_rank, leader_score, component_scores_json, coverage,
              lifecycle_status, leader_formula_version
       FROM l3_leader_pool_members
       WHERE pool_id=?
         AND lifecycle_status IN ('ACTIVE','NEW','REENTERED')
         AND leader_rank IS NOT NULL
         AND leader_rank <= 2
       ORDER BY leader_score DESC NULLS LAST, level1_name, level2_name, level3_name, leader_rank, stock_code""",
    (pool_id,),
)]

profiles = {
    r["stock_code"]: dict(r)
    for r in research.execute(
        """SELECT stock_code, main_business, main_products, business_scope, company_description
           FROM company_business_profiles"""
    )
}
research.close()

fundamentals: dict[str, dict] = {}
details: dict[str, dict] = {}
if TDX_DB.exists():
    tdx = sqlite3.connect(str(TDX_DB))
    fundamentals = load_payload_map(tdx, "fundamentals")
    details = load_payload_map(tdx, "security_details")
    tdx.close()

formula_version = leaders[0]["leader_formula_version"] if leaders else (
    pool_meta.get("formula_version") or "value-leader-v3.0.0"
)

rows: list[list[object]] = []
for index, row in enumerate(leaders, 1):
    status = "在池" if row["lifecycle_status"] == "ACTIVE" else row["lifecycle_status"]
    rows.append([
        index,
        row["stock_code"],
        row["stock_name"],
        row["level1_name"],
        row["level2_name"],
        row["level3_name"],
        row["leader_rank"],
        None if row["leader_score"] is None else float(row["leader_score"]),
        score_explanation(row),
        None if row["coverage"] is None else float(row["coverage"]),
        status,
        company_introduction(row["stock_code"], fundamentals, details, profiles),
    ])

unique_company_count = len({row[1] for row in rows})
unique_track_count = len({row[5] for row in rows})
introduction_available = sum(1 for row in rows if row[11] != "本地暂无可靠公司简介")

wb = Workbook()
ws = wb.active
ws.title = "龙头列表"
ws.sheet_view.showGridLines = False

title_fill = PatternFill("solid", fgColor="17365D")
meta_fill = PatternFill("solid", fgColor="EAF1FB")
note_fill = PatternFill("solid", fgColor="F8FAFD")
score_fill = PatternFill("solid", fgColor="FFF8E7")
header_fill = PatternFill("solid", fgColor="244062")
stripe_fill = PatternFill("solid", fgColor="F7FAFF")
thin = Border(
    left=Side(style="thin", color="D9E2F3"),
    right=Side(style="thin", color="D9E2F3"),
    top=Side(style="thin", color="D9E2F3"),
    bottom=Side(style="thin", color="D9E2F3"),
)

ws.merge_cells("A1:L1")
ws["A1"] = f"当前三级行业龙头池（{as_of}）"
ws["A1"].fill = title_fill
ws["A1"].font = Font(bold=True, color="FFFFFF", size=16)
ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[1].height = 30

ws.append([
    "数据口径：", f"当前 L3 Leader Pool 快照；数据日期 {as_of}",
    "细分赛道：", unique_track_count,
    "龙头关系：", len(rows),
    "去重公司：", unique_company_count,
    "简介可用：", introduction_available,
    "公式版本：", formula_version,
])
for col in range(1, 13):
    cell = ws.cell(2, col)
    cell.fill = meta_fill
    cell.font = Font(bold=(col % 2 == 1), color="244062", size=10)
ws.row_dimensions[2].height = 21

ws.merge_cells("A3:L3")
ws["A3"] = "说明：每一行代表一只公司在一个三级/末级细分赛道中的龙头席位；同一公司如覆盖多个赛道，会保留多行。按质量评分从高到低排序。"
ws["A3"].fill = note_fill
ws["A3"].font = Font(italic=True, color="5B6573", size=10)
ws["A3"].alignment = Alignment(wrap_text=True, vertical="center")
ws.row_dimensions[3].height = 30

ws.merge_cells("A4:L4")
ws["A4"] = (
    "评分说明：V3 两段式选龙头——赛道内排名由规模地位（市值/营收/净利）决定前2；"
    "综合评分（满分100）为质量分，由盈利质量、增长稳定性、现金流、估值与治理风险构成；"
    "规模地位单独展示在评分说明中。不是买卖建议。"
)
ws["A4"].fill = score_fill
ws["A4"].font = Font(italic=True, color="7A4E00", size=10)
ws["A4"].alignment = Alignment(wrap_text=True, vertical="center")
ws.row_dimensions[4].height = 36

ws.append(HEADER)
for col in range(1, 13):
    cell = ws.cell(5, col)
    cell.fill = header_fill
    cell.font = Font(bold=True, color="FFFFFF", size=10)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
ws.row_dimensions[5].height = 28

for row in rows:
    ws.append(row)

first_data = 6
last_data = 5 + len(rows)
for r in range(first_data, last_data + 1):
    for c in range(1, 13):
        cell = ws.cell(r, c)
        cell.font = Font(color="1F2937", size=10)
        cell.alignment = Alignment(
            vertical="center",
            wrap_text=(c in {9, 12}),
            horizontal="center" if c in {1, 2, 7, 8, 10, 11} else "left",
        )
        cell.border = thin
    ws.row_dimensions[r].height = 34
    ws.cell(r, 8).number_format = "0.00"
    ws.cell(r, 10).number_format = "0.0%"

ws.conditional_formatting.add(
    f"A{first_data}:L{last_data}",
    FormulaRule(formula=[f"MOD(ROW(),2)=0"], fill=stripe_fill),
)
ws.conditional_formatting.add(
    f"H{first_data}:H{last_data}",
    ColorScaleRule(
        start_type="min", start_color="FEE2E2",
        mid_type="percentile", mid_value=50, mid_color="FEF3C7",
        end_type="max", end_color="DCFCE7",
    ),
)

widths = {
    "A": 7, "B": 14, "C": 16, "D": 15, "E": 18, "F": 24,
    "G": 15, "H": 12, "I": 58, "J": 10, "K": 10, "L": 46,
}
for col, width in widths.items():
    ws.column_dimensions[col].width = width
ws.freeze_panes = "A6"
ws.auto_filter.ref = f"A5:L{last_data}"

xlsx_path = OUT_DIR / f"当前三级行业龙头池_{as_of}.xlsx"
wb.save(xlsx_path)

# Browser-friendly HTML with the same columns (user has no Excel/WPS)
html_path = OUT_DIR / f"当前三级行业龙头池_{as_of}.html"
html_rows = "".join(
    "<tr>" + "".join(f"<td>{'' if v is None else v}</td>" for v in row) + "</tr>"
    for row in rows
)
html = f"""<!DOCTYPE html>
<html lang="zh-CN"><head><meta charset="utf-8">
<title>当前三级行业龙头池（{as_of}）</title>
<style>
body{{font-family:'Segoe UI','Microsoft YaHei',sans-serif;margin:24px;background:#f7f7f5;color:#1a1a1a}}
h1{{font-size:22px;margin:0 0 8px}}.meta{{color:#444;margin:8px 0 16px;line-height:1.6}}
.wrap{{max-height:78vh;overflow:auto;background:#fff;border:1px solid #e5e5e2;border-radius:10px}}
table{{border-collapse:collapse;width:100%;font-size:12px}}
th,td{{border-bottom:1px solid #eee;padding:6px 8px;text-align:left;vertical-align:top}}
th{{position:sticky;top:0;background:#244062;color:#fff;z-index:1}}
td:nth-child(1),td:nth-child(2),td:nth-child(7),td:nth-child(8),td:nth-child(10),td:nth-child(11){{text-align:center;white-space:nowrap}}
td:nth-child(9),td:nth-child(12){{min-width:220px}}
tr:nth-child(even){{background:#f7faff}}
</style></head><body>
<h1>当前三级行业龙头池（{as_of}）</h1>
<div class="meta">
数据口径：当前 L3 Leader Pool 快照；细分赛道 {unique_track_count}；龙头关系 {len(rows)}；
去重公司 {unique_company_count}；简介可用 {introduction_available}；公式版本 {formula_version}
</div>
<div class="wrap"><table><thead><tr>{''.join(f'<th>{h}</th>' for h in HEADER)}</tr></thead>
<tbody>{html_rows}</tbody></table></div>
</body></html>"""
html_path.write_text(html, encoding="utf-8")

print(json.dumps({
    "xlsx": str(xlsx_path),
    "html": str(html_path),
    "as_of": as_of,
    "pool_id": pool_id,
    "formula_version": formula_version,
    "leaders": len(rows),
    "tracks": unique_track_count,
    "companies": unique_company_count,
    "intro_available": introduction_available,
}, ensure_ascii=False, indent=2))
